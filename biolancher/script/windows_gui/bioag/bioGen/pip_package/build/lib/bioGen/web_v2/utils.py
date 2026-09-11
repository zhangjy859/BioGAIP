import os
import sys
import time
import signal
import logging
import collections
from typing import Deque, Tuple
import pypdf
from autogen_core.logging import LLMCallEvent
import streamlit as st

class LLMUsageTracker(logging.Handler):
    def __init__(self) -> None:
        super().__init__()
        self._prompt_tokens = 0
        self._completion_tokens = 0

    @property
    def tokens(self) -> int: return self._prompt_tokens + self._completion_tokens

    @property
    def prompt_tokens(self) -> int: return self._prompt_tokens

    @property
    def completion_tokens(self) -> int: return self._completion_tokens

    def reset(self) -> None:
        self._prompt_tokens = 0
        self._completion_tokens = 0

    def emit(self, record: logging.LogRecord) -> None:
        try:
            if isinstance(record.msg, LLMCallEvent):
                self._prompt_tokens += record.msg.prompt_tokens
                self._completion_tokens += record.msg.completion_tokens
        except Exception:
            self.handleError(record)

def restart_session():
    os.execv(sys.executable, [sys.executable] + sys.argv)

agent_avatars = {
    "user": "🧑", "user_proxy": "👤", "system_check_agent": "🔍",
    "bio_micromamba_env_agent": "🧪", "file_agent": "📁", "web_surfer_agent": "🌐",
    "manager_mem_agent": "🧠", "system": "⚙️", "planning_agent": "🗂️",
    "query_agent": "❓", "command_generator_agent": "💻", "code_generator_agent": "📝",
    "safety_checker_agent": "🛡️", "executor_agent": "🚀", "assistant": "🤖",
}

def get_avatar(name):
    return agent_avatars.get(name, "🤖")

def force_kill_process(pid_info, logger=None):
    pid = None
    pid_file = None
    if isinstance(pid_info, str):
        if os.path.exists(pid_info):
            pid_file = pid_info
            with open(pid_info, 'r') as f:
                pid = int(f.read().strip())
        else:
            pid = int(pid_info)
    else:
        pid = int(pid_info)
    try:
        os.kill(pid, signal.SIGKILL)
        time.sleep(3)
        try: os.waitpid(pid, os.WNOHANG)
        except Exception: pass
        if pid_file and os.path.exists(pid_file):
            os.remove(pid_file)
    except Exception as e:
        if logger: logger.error(f'Failed to force kill process {pid}: {e}')

def read_uploaded_file(uploaded_file):
    file_name = uploaded_file.name
    file_ext = file_name.split('.')[-1].lower() if '.' in file_name else 'txt'
    try:
        if uploaded_file.type == "application/pdf":
            reader = pypdf.PdfReader(uploaded_file)
            text = "".join([page.extract_text() + "\n" for page in reader.pages])
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            import openpyxl
            wb = openpyxl.load_workbook(uploaded_file)
            text = ""
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                text += f"Sheet: {sheet_name}\n"
                for row in ws.iter_rows(values_only=True):
                    text += "\t".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
                text += "\n"
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            import docx
            doc = docx.Document(uploaded_file)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            import pptx
            prs = pptx.Presentation(uploaded_file)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"): text += shape.text + "\n"
        else:
            text = uploaded_file.read().decode('utf-8')
        return f'<uploaded_file name="{file_name}" ext="{file_ext}">\n{text.strip()}\n</uploaded_file>'
    except Exception as e:
        st.error(f"Error reading file: {str(e)}")
        return None

class Reward:
    def __init__(self, initial_reward: float, decay_rate: float = 0.6, slow_decay_rate: float = 0.99, recovery_step: float = -1, cooldown_period: int = 10, window_a: int = 15, min_triggers_a: int = 3, window_b: int = 30, min_pattern_len: int = 3):
        self.initial_reward = initial_reward
        self.decay_rate = decay_rate
        self.slow_decay_rate = slow_decay_rate
        self.recovery_step = 0.1 if recovery_step < 0 else recovery_step
        self.cooldown_period = cooldown_period
        self.window_a = window_a
        self.min_triggers_a = min_triggers_a
        self.window_b = window_b
        self.min_pattern_len = min_pattern_len
        self.reset()

    def reset(self) -> None:
        self.current_reward = float(self.initial_reward)
        self.clock = 0
        self.history: Deque[Tuple[int, str]] = collections.deque()
        self.cooldown_until = 0

    def timer(self) -> None:
        self.clock += 1

    def __call__(self, category: str = 'default') -> int:
        self.history.append((self.clock, category))
        self._maintain_history()
        self._evaluate_recovery()
        output_reward = round(self.current_reward)
        if output_reward == 0:
            self.current_reward *= self.slow_decay_rate
        else:
            self.current_reward *= self.decay_rate
        return output_reward

    def _maintain_history(self) -> None:
        while self.history and self.history[0][0] <= self.clock - self.window_b:
            self.history.popleft()

    def _evaluate_recovery(self) -> None:
        if self.clock <= self.cooldown_until: return
        recent_triggers_a = sum(1 for tick, _ in self.history if tick > self.clock - self.window_a)
        if recent_triggers_a < self.min_triggers_a: return
        recent_b = list(self.history)
        max_len = 0
        n = len(recent_b)
        for i in range(n - self.min_pattern_len + 1):
            cat = recent_b[i][1]
            if recent_b[i + 1][1] != cat: continue
            interval = recent_b[i + 1][0] - recent_b[i][0]
            curr_len = 2
            for j in range(i + 2, n):
                if recent_b[j][1] == cat and recent_b[j][0] - recent_b[j - 1][0] == interval:
                    curr_len += 1
                else: break
            if curr_len >= self.min_pattern_len and recent_b[i + curr_len - 1][0] == self.clock:
                max_len = max(max_len, curr_len)
        if max_len >= self.min_pattern_len:
            recovery_amount = self.initial_reward * (max_len * self.recovery_step)
            self.current_reward = min(self.initial_reward, self.current_reward + recovery_amount)
            self.cooldown_until = self.clock + self.cooldown_period